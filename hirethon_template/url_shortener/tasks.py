import openpyxl
from io import BytesIO
from celery import shared_task
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from .models import BulkUploadTask, ShortURL, Namespace


@shared_task(bind=True)
def process_bulk_upload_task(self, task_id):
    """
    Process bulk URL upload from Excel file.
    Expected columns: original_url, custom_short_code (optional)
    """
    try:
        task = BulkUploadTask.objects.get(id=task_id)
        task.status = BulkUploadTask.Status.PROCESSING
        task.save()
        
        # Read the input Excel file
        input_file = task.input_file
        wb = openpyxl.load_workbook(input_file)
        ws = wb.active
        
        # Get headers from first row - normalize them
        raw_headers = [cell.value for cell in ws[1]]
        # Strip whitespace and convert to lowercase for comparison
        headers = [str(h).strip().lower() if h else '' for h in raw_headers]
        
        # Find column indices
        try:
            url_col_idx = headers.index('original_url')
        except ValueError:
            task.status = BulkUploadTask.Status.FAILED
            task.error_log = f"Required column 'original_url' not found in Excel file. Found columns: {', '.join(raw_headers)}"
            task.save()
            return
        
        # Check for optional custom_short_code column
        try:
            code_col_idx = headers.index('custom_short_code')
            has_custom_code = True
        except ValueError:
            has_custom_code = False
            code_col_idx = None
        
        # Prepare output workbook
        output_wb = openpyxl.Workbook()
        output_ws = output_wb.active
        output_ws.append(['original_url', 'short_code', 'short_url', 'status', 'error'])
        
        total_urls = ws.max_row - 1  # Exclude header
        task.total_urls = total_urls
        task.save()
        
        processed = 0
        failed = 0
        error_messages = []
        
        # Process each row (skip header)
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                original_url = row[url_col_idx]
                custom_short_code = row[code_col_idx] if has_custom_code and code_col_idx < len(row) else None
                
                if not original_url:
                    error_messages.append(f"Row {row_idx}: Empty URL")
                    output_ws.append([original_url, '', '', 'FAILED', 'Empty URL'])
                    failed += 1
                    continue
                
                # Create ShortURL
                short_url_data = {
                    'namespace': task.namespace,
                    'original_url': str(original_url),
                    'created_by': task.created_by
                }
                
                if custom_short_code:
                    short_code_to_use = str(custom_short_code).strip()
                    # Check if short code already exists
                    exists = ShortURL.objects.filter(
                        namespace=task.namespace,
                        short_code=short_code_to_use
                    ).exists()
                    if exists:
                        error_messages.append(f"Row {row_idx}: Short code '{short_code_to_use}' already exists")
                        output_ws.append([original_url, short_code_to_use, '', 'FAILED', 'Short code already exists'])
                        failed += 1
                        continue
                    short_url_data['short_code'] = short_code_to_use
                else:
                    # Auto-generate short code
                    from .models import generate_short_code
                    max_attempts = 10
                    for _ in range(max_attempts):
                        generated_code = generate_short_code()
                        if not ShortURL.objects.filter(namespace=task.namespace, short_code=generated_code).exists():
                            short_url_data['short_code'] = generated_code
                            break
                    else:
                        error_messages.append(f"Row {row_idx}: Unable to generate unique short code")
                        output_ws.append([original_url, '', '', 'FAILED', 'Unable to generate unique short code'])
                        failed += 1
                        continue
                
                short_url = ShortURL.objects.create(**short_url_data)
                
                # Generate full absolute URL
                full_url = short_url.get_full_short_url()
                
                output_ws.append([
                    original_url,
                    short_url.short_code,
                    full_url,
                    'SUCCESS',
                    ''
                ])
                processed += 1
                
            except Exception as e:
                error_msg = str(e)
                error_messages.append(f"Row {row_idx}: {error_msg}")
                output_ws.append([
                    original_url if 'original_url' in locals() else '',
                    custom_short_code if 'custom_short_code' in locals() else '',
                    '',
                    'FAILED',
                    error_msg
                ])
                failed += 1
            
            # Update progress
            task.processed_urls = processed + failed
            task.failed_urls = failed
            task.save()
        
        # Save output Excel file
        output_buffer = BytesIO()
        output_wb.save(output_buffer)
        output_buffer.seek(0)
        
        output_filename = f"bulk_upload_result_{task_id}.xlsx"
        
        # Use the configured storage (S3 if enabled) with fallback to local
        try:
            # Use the same storage approach as QR code generation
            from hirethon_template.utils.storages import MediaRootS3Boto3Storage
            s3_storage = MediaRootS3Boto3Storage()
            file_path = s3_storage.save(f"bulk_uploads/output/{output_filename}", ContentFile(output_buffer.read()))
            task.output_file = file_path
            task.save()
        except Exception as e:
            # Fallback to local storage if S3 fails
            output_buffer.seek(0)
            from django.core.files.storage import FileSystemStorage
            local_storage = FileSystemStorage()
            file_path = local_storage.save(f"bulk_uploads/output/{output_filename}", ContentFile(output_buffer.read()))
            task.output_file = file_path
            task.save()
        
        # Update task status
        task.status = BulkUploadTask.Status.COMPLETED
        task.error_log = '\n'.join(error_messages) if error_messages else ''
        task.save()
        
        return {
            'task_id': task_id,
            'status': 'completed',
            'processed': processed,
            'failed': failed
        }
        
    except BulkUploadTask.DoesNotExist:
        return {'error': f'Task {task_id} not found'}
    except Exception as e:
        # Mark task as failed
        try:
            task = BulkUploadTask.objects.get(id=task_id)
            task.status = BulkUploadTask.Status.FAILED
            task.error_log = f"Unexpected error: {str(e)}"
            task.save()
        except:
            pass
        raise

